"""Builds excel/Horse_Scoring_Bloodline_Analytics.xlsx — every analysis cell is a live formula."""
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, ScatterChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

SRC, OUT = sys.argv[1], sys.argv[2]
profile = pd.read_excel(SRC, sheet_name='Horse_Profile')
scores = pd.read_excel(SRC, sheet_name='Judge_Scores')
parents = pd.read_excel(SRC, sheet_name='Parent_Reference')
TRAITS = ['Type', 'Head_Neck', 'Body_Topline', 'Legs', 'Movement', 'Balance']
NH, NS, NP = len(profile), len(scores), len(parents)          # 40, 120, 15
HL, SL, PL = NH + 1, NS + 1, NP + 1                            # last data rows

# Ordering for analysis tables (static labels; all numbers are formulas)
T = scores.copy(); T['Tot'] = T[TRAITS].sum(axis=1)
hav = T.groupby('HorseID').Tot.mean()
prof = profile.set_index('HorseID').assign(Avg=hav)
SIRES = prof.groupby('Sire').Avg.mean().sort_values(ascending=False).index.tolist()
DAMSIRES = prof.groupby('DamSire').Avg.mean().sort_values(ascending=False).index.tolist()
DAMLINES = prof.groupby('DamLine').Avg.mean().sort_values(ascending=False).index.tolist()
FARMS = prof.groupby('FarmGroup').Avg.mean().sort_values(ascending=False).index.tolist()
AGES = ['Junior', 'Young Horse', 'Mature']
SEXES = ['Colt', 'Filly', 'Gelding', 'Mare', 'Stallion']
JUDGES = sorted(scores.Judge.unique())

# ---------- styles ----------
NAVY, BLUE, LIGHT, ZEBRA, LINE = '1F3A5F', '2A78D6', 'E8F1FC', 'F6F8FB', 'D5DCE6'
F = lambda **k: Font(name='Arial', **{'size': 10, **k})
HDR_FILL = PatternFill('solid', fgColor=NAVY)
SUB_FILL = PatternFill('solid', fgColor=LIGHT)
INPUT_FILL = PatternFill('solid', fgColor='FFF2CC')
thin = Side(style='thin', color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

wb = Workbook()
wb.remove(wb.active)


def sheet(name, tab=None):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    if tab: ws.sheet_properties.tabColor = tab
    return ws


def title(ws, text, sub=None, width_to='L'):
    ws['A1'] = text; ws['A1'].font = F(size=16, bold=True, color=NAVY)
    if sub: ws['A2'] = sub; ws['A2'].font = F(size=10, italic=True, color='52514E')
    ws.row_dimensions[1].height = 26


def header(ws, row, col, labels, fill=HDR_FILL, color='FFFFFF'):
    for i, h in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=h)
        c.font = F(bold=True, color=color); c.fill = fill; c.alignment = CENTER; c.border = BOX


def section(ws, cell, text):
    ws[cell] = text; ws[cell].font = F(size=12, bold=True, color=NAVY)


def body(ws, r1, r2, c1, c2, fmt=None, zebra=True):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            from copy import copy as _c; fnt=_c(cell.font); fnt.name='Arial'; cell.font=fnt
            cell.border = BOX
            if zebra and (r - r1) % 2 == 1: cell.fill = PatternFill('solid', fgColor=ZEBRA)
            if fmt and c in fmt: cell.number_format = fmt[c]


def widths(ws, spec):
    for col, w in spec.items(): ws.column_dimensions[col].width = w


def color_bars(chart, hexcolor):
    for s in chart.series:
        s.graphicalProperties.solidFill = hexcolor
        s.graphicalProperties.line.solidFill = hexcolor


def style_chart(ch, title_txt, w=16, h=7.5):
    ch.title = title_txt; ch.width = w; ch.height = h; ch.style = 10
    ch.legend = None
    return ch


div_scale = lambda: ColorScaleRule(start_type='num', start_value=-1.2, start_color='E34948',
                                   mid_type='num', mid_value=0, mid_color='F7F6F3',
                                   end_type='num', end_value=1.2, end_color='2A78D6')
seq_scale = lambda lo, hi: ColorScaleRule(start_type='num', start_value=lo, start_color='EAF2FD',
                                          end_type='num', end_value=hi, end_color='2A78D6')

# ======================= Settings =======================
st = sheet('Settings', '8A8984')
title(st, 'Settings & assumptions', 'Yellow cells with blue text are inputs — change them and every sheet recalculates.')
rows = [('Show / reference year', 2026, 'Used to compute age from BirthYear. Assumed current season.'),
        ('Max score per trait', 20, 'AHA-style halter category maximum (from dataset README).'),
        ('Max total per judge', 120, '6 categories x 20 points (from dataset README).'),
        ('Elite tier — minimum avg total', 112, 'Tier cut-offs carried over from the source workbook formulas.'),
        ('High tier — minimum avg total', 106, ''),
        ('Developing tier — minimum avg total', 100, 'Below this = "Needs Review".'),
        ('Colt/Filly maximum age', 3, 'Halter convention: colt/filly under 4, stallion/mare 4+. Used for data-quality flag.')]
header(st, 4, 1, ['Setting', 'Value', 'Note / source'])
for i, (a, b, c) in enumerate(rows, start=5):
    st.cell(row=i, column=1, value=a).font = F()
    v = st.cell(row=i, column=2, value=b); v.font = F(color='0000FF', bold=True); v.fill = INPUT_FILL
    st.cell(row=i, column=3, value=c).font = F(color='52514E')
body(st, 5, 11, 1, 3, zebra=False)
for r in range(5, 12): st.cell(row=r, column=2).fill = INPUT_FILL
widths(st, {'A': 36, 'B': 12, 'C': 80})
S_YEAR, S_ELITE, S_HIGH, S_DEV, S_COLTMAX = 'Settings!$B$5', 'Settings!$B$8', 'Settings!$B$9', 'Settings!$B$10', 'Settings!$B$11'

# ======================= Horse_Profile =======================
hp = sheet('Horse_Profile', '1F3A5F')
cols = list(profile.columns) + ['Age', 'DataFlag']
header(hp, 1, 1, cols)
for r, row in enumerate(profile.itertuples(index=False), start=2):
    for c, v in enumerate(row, start=1):
        hp.cell(row=r, column=c, value=v)
    hp.cell(row=r, column=16, value=f'={S_YEAR}-E{r}')
    hp.cell(row=r, column=17, value=(f'=IF(AND(OR(D{r}="Colt",D{r}="Filly"),P{r}>{S_COLTMAX}),"Sex/age mismatch",'
                                      f'IF(AND(OR(D{r}="Stallion",D{r}="Mare"),P{r}<={S_COLTMAX}),"Sex/age mismatch",""))'))
body(hp, 2, HL, 1, 17, zebra=False)
hp.add_table(Table(displayName='tblHorseProfile', ref=f'A1:Q{HL}',
                   tableStyleInfo=TableStyleInfo(name='TableStyleLight9', showRowStripes=True)))
hp.freeze_panes = 'C2'
widths(hp, {L(i): w for i, w in enumerate([9, 26, 11, 10, 10, 13, 10, 14, 22, 20, 18, 24, 20, 20, 24, 7, 17], start=1)})
hp['Q1'].comment = Comment('Flags sex labels that conflict with age (colt/filly must be under 4; stallion/mare 4+). Threshold on Settings.', 'Analyst')

# ======================= Judge_Scores =======================
js = sheet('Judge_Scores', '1F3A5F')
header(js, 1, 1, list(scores.columns))
for r, row in enumerate(scores.itertuples(index=False), start=2):
    for c, v in enumerate(row, start=1):
        js.cell(row=r, column=c, value=v)
    js.cell(row=r, column=11, value=f'=SUM(E{r}:J{r})')
body(js, 2, SL, 1, 12, fmt={k: '0.0' for k in range(5, 12)}, zebra=False)
js.add_table(Table(displayName='tblJudgeScores', ref=f'A1:L{SL}',
                   tableStyleInfo=TableStyleInfo(name='TableStyleLight9', showRowStripes=True)))
js.freeze_panes = 'E2'
widths(js, {'A': 10, 'B': 9, 'C': 26, 'D': 10, **{L(i): 11 for i in range(5, 12)}, 'L': 52})

# ======================= Parent_Reference =======================
pr = sheet('Parent_Reference', '1F3A5F')
header(pr, 1, 1, list(parents.columns) + ['Offspring', 'OffspringAvg', 'OffspringVsParent'])
for r, row in enumerate(parents.itertuples(index=False), start=2):
    for c, v in enumerate(row, start=1):
        pr.cell(row=r, column=c, value=v)
    # Offspring count / avg: sires match Horse_Summary!C, dam sires match Horse_Summary!D
    pr.cell(row=r, column=8, value=f'=IF(B{r}="Sire",COUNTIF(Horse_Summary!$C$2:$C${HL},A{r}),COUNTIF(Horse_Summary!$D$2:$D${HL},A{r}))')
    pr.cell(row=r, column=9, value=(f'=IFERROR(IF(B{r}="Sire",AVERAGEIF(Horse_Summary!$C$2:$C${HL},A{r},Horse_Summary!$O$2:$O${HL}),'
                                     f'AVERAGEIF(Horse_Summary!$D$2:$D${HL},A{r},Horse_Summary!$O$2:$O${HL})),"")'))
    pr.cell(row=r, column=10, value=f'=IFERROR(I{r}-D{r},"")')
body(pr, 2, PL, 1, 10, fmt={4: '0.0', 9: '0.00', 10: '+0.00;-0.00;0.00'}, zebra=False)
pr.add_table(Table(displayName='tblParents', ref=f'A1:J{PL}',
                   tableStyleInfo=TableStyleInfo(name='TableStyleLight9', showRowStripes=True)))
widths(pr, {'A': 22, 'B': 11, 'C': 24, 'D': 13, 'E': 13, 'F': 15, 'G': 50, 'H': 10, 'I': 13, 'J': 16})

# ======================= Horse_Summary =======================
hs = sheet('Horse_Summary', '2A78D6')
HS_COLS = ['HorseID', 'RegisteredName', 'Sire', 'DamSire', 'Sex', 'AgeGroup', 'PaternalLine', 'DamLine'] + TRAITS + \
          ['AvgTotal', 'JudgeSpread', 'Rank', 'UniqueRank', 'ScoreTier', 'MainStrength', 'FocusArea',
           'SireScore', 'DamSireScore', 'MidParentScore', 'RankInSireLine', 'SireKey', 'VsSireScore', 'DataFlag']
header(hs, 1, 1, HS_COLS)
pcol = {'RegisteredName': 'B', 'Sire': 'I', 'DamSire': 'K', 'Sex': 'D', 'AgeGroup': 'F', 'PaternalLine': 'L', 'DamLine': 'M'}
JR = f'Judge_Scores!$B$2:$B${SL}'
for i, hid in enumerate(profile.HorseID, start=2):
    r = i
    hs.cell(row=r, column=1, value=hid)
    for j, name in enumerate(['RegisteredName', 'Sire', 'DamSire', 'Sex', 'AgeGroup', 'PaternalLine', 'DamLine'], start=2):
        c = pcol[name]
        hs.cell(row=r, column=j, value=f'=INDEX(Horse_Profile!${c}$2:${c}${HL},MATCH($A{r},Horse_Profile!$A$2:$A${HL},0))')
    for k, src_col in enumerate('EFGHIJ'):
        hs.cell(row=r, column=9 + k, value=f'=AVERAGEIF({JR},$A{r},Judge_Scores!${src_col}$2:${src_col}${SL})')
    f = {
        'O': f'=AVERAGEIF({JR},$A{r},Judge_Scores!$K$2:$K${SL})',
        'P': f'=_xlfn.MAXIFS(Judge_Scores!$K$2:$K${SL},{JR},$A{r})-_xlfn.MINIFS(Judge_Scores!$K$2:$K${SL},{JR},$A{r})',
        'Q': f'=RANK(O{r},$O$2:$O${HL},0)',
        'R': f'=Q{r}+COUNTIF($O$2:O{r},O{r})-1',
        'S': f'=IF(O{r}>={S_ELITE},"Elite",IF(O{r}>={S_HIGH},"High",IF(O{r}>={S_DEV},"Developing","Needs Review")))',
        'T': f'=INDEX($I$1:$N$1,MATCH(MAX(I{r}:N{r}),I{r}:N{r},0))',
        'U': f'=INDEX($I$1:$N$1,MATCH(MIN(I{r}:N{r}),I{r}:N{r},0))',
        'V': f'=INDEX(Parent_Reference!$D$2:$D${PL},MATCH(C{r},Parent_Reference!$A$2:$A${PL},0))',
        'W': f'=INDEX(Parent_Reference!$D$2:$D${PL},MATCH(D{r},Parent_Reference!$A$2:$A${PL},0))',
        'X': f'=AVERAGE(V{r}:W{r})',
        'Y': f'=COUNTIFS($C$2:$C${HL},C{r},$O$2:$O${HL},">"&O{r})+1',
        'Z': f'=C{r}&"|"&Y{r}',
        'AA': f'=O{r}-V{r}',
        'AB': f'=INDEX(Horse_Profile!$Q$2:$Q${HL},MATCH($A{r},Horse_Profile!$A$2:$A${HL},0))',
    }
    for col, formula in f.items():
        hs[f'{col}{r}'] = formula
fmt = {k: '0.00' for k in range(9, 17)}; fmt.update({22: '0.0', 23: '0.0', 24: '0.00', 27: '+0.00;-0.00;0.00'})
body(hs, 2, HL, 1, 28, fmt=fmt, zebra=False)
hs.add_table(Table(displayName='tblHorseSummary', ref=f'A1:AB{HL}',
                   tableStyleInfo=TableStyleInfo(name='TableStyleLight9', showRowStripes=True)))
hs.freeze_panes = 'C2'
hs.conditional_formatting.add(f'O2:O{HL}', seq_scale(104, 114))
hs.conditional_formatting.add(f'S2:S{HL}', CellIsRule(operator='equal', formula=['"Elite"'], font=Font(name='Arial', bold=True, color='1C5CAB')))
hs.conditional_formatting.add(f'AB2:AB{HL}', CellIsRule(operator='notEqual', formula=['""'], fill=PatternFill('solid', fgColor='FDE2E1')))
widths(hs, {'A': 9, 'B': 26, 'C': 21, 'D': 17, 'E': 9, 'F': 12, 'G': 22, 'H': 19, **{L(i): 11 for i in range(9, 17)},
            'Q': 7, 'R': 8, 'S': 13, 'T': 14, 'U': 14, 'V': 10, 'W': 12, 'X': 13, 'Y': 12, 'Z': 24, 'AA': 12, 'AB': 17})
hs['R1'].comment = Comment('Rank with ties broken by row order — used to build the Top 10 list.', 'Analyst')
hs['Z1'].comment = Comment('Helper key (Sire|rank within sire line) used to look up each line\'s top horse.', 'Analyst')
hs['P1'].comment = Comment('Highest minus lowest judge total for this horse — a measure of judge disagreement.', 'Analyst')

HSR = lambda col: f'Horse_Summary!${col}$2:${col}${HL}'
TRAIT_COL = dict(zip(TRAITS, 'IJKLMN'))

# ======================= Bloodline_Analysis =======================
ba = sheet('Bloodline_Analysis', 'EB6834')
title(ba, 'Bloodline analysis', 'Offspring performance by sire line, dam sire, dam line and farm group. All figures are live formulas on Horse_Summary.')
section(ba, 'A4', '1 · Sire line comparison')
h = ['Sire', 'Paternal line', 'Horses', 'Avg total', 'Std dev', 'Min', 'Max', 'Elite horses', 'Elite %',
     'Sire show score', 'Offspring vs sire', 'Top horse', 'Rank']
header(ba, 5, 1, h)
r0 = 6
for i, s in enumerate(SIRES):
    r = r0 + i
    ba[f'A{r}'] = s
    ba[f'B{r}'] = f'=INDEX(Parent_Reference!$C$2:$C${PL},MATCH(A{r},Parent_Reference!$A$2:$A${PL},0))'
    ba[f'C{r}'] = f'=COUNTIF({HSR("C")},A{r})'
    ba[f'D{r}'] = f'=AVERAGEIF({HSR("C")},A{r},{HSR("O")})'
    ba[f'E{r}'] = f'=SQRT(SUMPRODUCT(({HSR("C")}=A{r})*({HSR("O")}-D{r})^2)/(C{r}-1))'
    ba[f'F{r}'] = f'=_xlfn.MINIFS({HSR("O")},{HSR("C")},A{r})'
    ba[f'G{r}'] = f'=_xlfn.MAXIFS({HSR("O")},{HSR("C")},A{r})'
    ba[f'H{r}'] = f'=COUNTIFS({HSR("C")},A{r},{HSR("S")},"Elite")'
    ba[f'I{r}'] = f'=H{r}/C{r}'
    ba[f'J{r}'] = f'=INDEX(Parent_Reference!$D$2:$D${PL},MATCH(A{r},Parent_Reference!$A$2:$A${PL},0))'
    ba[f'K{r}'] = f'=D{r}-J{r}'
    ba[f'L{r}'] = f'=INDEX({HSR("B")},MATCH(A{r}&"|1",{HSR("Z")},0))'
    ba[f'M{r}'] = f'=RANK(D{r},$D${r0}:$D${r0 + len(SIRES) - 1},0)'
rS2 = r0 + len(SIRES) - 1
ba[f'A{rS2 + 1}'] = 'All horses'
ba[f'C{rS2 + 1}'] = f'=SUM(C{r0}:C{rS2})'
ba[f'D{rS2 + 1}'] = f'=AVERAGE({HSR("O")})'
ba[f'E{rS2 + 1}'] = f'=STDEV({HSR("O")})'
ba[f'F{rS2 + 1}'] = f'=MIN({HSR("O")})'
ba[f'G{rS2 + 1}'] = f'=MAX({HSR("O")})'
ba[f'H{rS2 + 1}'] = f'=SUM(H{r0}:H{rS2})'
ba[f'I{rS2 + 1}'] = f'=H{rS2 + 1}/C{rS2 + 1}'
body(ba, r0, rS2 + 1, 1, 13, fmt={4: '0.00', 5: '0.00', 6: '0.00', 7: '0.00', 9: '0%', 10: '0.0', 11: '+0.00;-0.00;0.00'})
for c in range(1, 14): ba.cell(row=rS2 + 1, column=c).font = F(bold=True); ba.cell(row=rS2 + 1, column=c).fill = SUB_FILL
ba.conditional_formatting.add(f'D{r0}:D{rS2}', seq_scale(105, 113))

# Trait averages by sire
t0 = rS2 + 4
section(ba, f'A{t0 - 1}', '2 · Conformation trait averages by sire')
header(ba, t0, 1, ['Sire'] + TRAITS + ['Strongest', 'Weakest'])
for i, s in enumerate(SIRES + ['All horses']):
    r = t0 + 1 + i
    ba[f'A{r}'] = s
    for k, t in enumerate(TRAITS):
        c = TRAIT_COL[t]
        ba.cell(row=r, column=2 + k, value=(f'=AVERAGE({HSR(c)})' if s == 'All horses'
                                            else f'=AVERAGEIF({HSR("C")},$A{r},{HSR(c)})'))
    ba[f'H{r}'] = f'=INDEX($B${t0}:$G${t0},MATCH(MAX(B{r}:G{r}),B{r}:G{r},0))'
    ba[f'I{r}'] = f'=INDEX($B${t0}:$G${t0},MATCH(MIN(B{r}:G{r}),B{r}:G{r},0))'
tAll = t0 + 1 + len(SIRES)
body(ba, t0 + 1, tAll, 1, 9, fmt={k: '0.00' for k in range(2, 8)})
for c in range(1, 10): ba.cell(row=tAll, column=c).font = F(bold=True); ba.cell(row=tAll, column=c).fill = SUB_FILL

d0 = tAll + 3
section(ba, f'A{d0 - 1}', '3 · Difference from all-horse average (blue = strength, red = weakness)')
header(ba, d0, 1, ['Sire'] + TRAITS + ['Net (sum)'])
for i, s in enumerate(SIRES):
    r = d0 + 1 + i; src = t0 + 1 + i
    ba[f'A{r}'] = f'=A{src}'
    for k in range(6):
        cl = L(2 + k)
        ba[f'{cl}{r}'] = f'={cl}{src}-{cl}${tAll}'
    ba[f'H{r}'] = f'=SUM(B{r}:G{r})'
dE = d0 + len(SIRES)
body(ba, d0 + 1, dE, 1, 8, fmt={k: '+0.00;-0.00;0.00' for k in range(2, 9)}, zebra=False)
ba.conditional_formatting.add(f'B{d0 + 1}:G{dE}', div_scale())

# Dam sire, dam line, farm group
def group_table(start_row, label, items, key_col, extra_parent_score=False):
    section(ba, f'A{start_row - 1}', label)
    hd = ['Group', 'Horses', 'Avg total', 'Min', 'Max', 'Elite horses']
    if extra_parent_score: hd += ['Dam-sire show score', 'Offspring vs score']
    header(ba, start_row, 1, hd)
    for i, g in enumerate(items):
        r = start_row + 1 + i
        ba[f'A{r}'] = g
        ba[f'B{r}'] = f'=COUNTIF({HSR(key_col)},A{r})'
        ba[f'C{r}'] = f'=AVERAGEIF({HSR(key_col)},A{r},{HSR("O")})'
        ba[f'D{r}'] = f'=_xlfn.MINIFS({HSR("O")},{HSR(key_col)},A{r})'
        ba[f'E{r}'] = f'=_xlfn.MAXIFS({HSR("O")},{HSR(key_col)},A{r})'
        ba[f'F{r}'] = f'=COUNTIFS({HSR(key_col)},A{r},{HSR("S")},"Elite")'
        if extra_parent_score:
            ba[f'G{r}'] = f'=INDEX(Parent_Reference!$D$2:$D${PL},MATCH(A{r},Parent_Reference!$A$2:$A${PL},0))'
            ba[f'H{r}'] = f'=C{r}-G{r}'
    end = start_row + len(items)
    body(ba, start_row + 1, end, 1, 8 if extra_parent_score else 6,
         fmt={3: '0.00', 4: '0.00', 5: '0.00', 7: '0.0', 8: '+0.00;-0.00;0.00'})
    ba.conditional_formatting.add(f'C{start_row + 1}:C{end}', seq_scale(104, 113))
    return end

g1 = dE + 3
e1 = group_table(g1, '4 · Dam-sire comparison', DAMSIRES, 'D', True)
e2 = group_table(e1 + 3, '5 · Dam line comparison', DAMLINES, 'H')
farm_col = None
# Farm group lives on Horse_Profile only — use direct references there
section(ba, f'A{e2 + 2}', '6 · Farm group comparison')
g3 = e2 + 3
header(ba, g3, 1, ['Group', 'Horses', 'Avg total', 'Min', 'Max', 'Elite horses'])
for i, g in enumerate(FARMS):
    r = g3 + 1 + i
    ba[f'A{r}'] = g
    ba[f'B{r}'] = f'=COUNTIF(Horse_Profile!$N$2:$N${HL},A{r})'
    ba[f'C{r}'] = f'=SUMPRODUCT((Horse_Profile!$N$2:$N${HL}=A{r})*{HSR("O")})/B{r}'
    ba[f'D{r}'] = f'=_xlfn.MINIFS({HSR("O")},Horse_Profile!$N$2:$N${HL},A{r})'
    ba[f'E{r}'] = f'=_xlfn.MAXIFS({HSR("O")},Horse_Profile!$N$2:$N${HL},A{r})'
    ba[f'F{r}'] = f'=COUNTIFS(Horse_Profile!$N$2:$N${HL},A{r},{HSR("S")},"Elite")'
body(ba, g3 + 1, g3 + len(FARMS), 1, 6, fmt={3: '0.00', 4: '0.00', 5: '0.00'})
ba[f'A{g3 + len(FARMS) + 1}'] = 'Note: farm-group rows rely on Horse_Profile and Horse_Summary sharing the same HorseID order.'
ba[f'A{g3 + len(FARMS) + 1}'].font = F(size=9, italic=True, color='8A8984')

widths(ba, {'A': 24, 'B': 22, 'C': 11, 'D': 11, 'E': 11, 'F': 11, 'G': 13, 'H': 13, 'I': 12, 'J': 13, 'K': 14, 'L': 24, 'M': 7})

ch = style_chart(BarChart(), 'Average offspring score by sire', 15, 7.5)
ch.type = 'bar'; ch.add_data(Reference(ba, min_col=4, min_row=5, max_row=rS2), titles_from_data=True)
ch.set_categories(Reference(ba, min_col=1, min_row=r0, max_row=rS2))
ch.y_axis.scaling.min = 100; ch.y_axis.scaling.max = 115; ch.x_axis.scaling.orientation = 'maxMin'
ch.dataLabels = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False, showPercent=False); ch.dataLabels.numFmt = '0.0'
ch.y_axis.delete = False; ch.x_axis.delete = False; ch.y_axis.majorGridlines = None
color_bars(ch, BLUE); ba.add_chart(ch, 'O4')

ch = BarChart(); ch.type = 'col'; ch.grouping = 'clustered'
ch.add_data(Reference(ba, min_col=2, max_col=7, min_row=d0, max_row=dE), titles_from_data=True)
ch.set_categories(Reference(ba, min_col=1, min_row=d0 + 1, max_row=dE))
ch.title = 'Trait strength vs herd average, by sire'; ch.width = 18; ch.height = 8; ch.style = 10
ch.y_axis.title = 'Points vs average'; ch.y_axis.delete = False; ch.x_axis.delete = False
ch.y_axis.number_format = '+0.0;-0.0;0'; ch.x_axis.tickLblPos = 'low'
ch.legend.position = 'b'
for s, col in zip(ch.series, ['2A78D6', 'EB6834', '1BAF7A', 'EDA100', 'E87BA4', '4A3AA7']):
    s.graphicalProperties.solidFill = col; s.graphicalProperties.line.solidFill = col
ba.add_chart(ch, 'O21')

# ======================= Parent_Performance =======================
pp = sheet('Parent_Performance', 'EB6834')
title(pp, 'Parent performance vs offspring', 'Do parent show scores predict how the offspring score? Correlations use horse-level averages from Horse_Summary.')
section(pp, 'A4', '1 · Predictive strength')
header(pp, 5, 1, ['Comparison', 'Pairs (n)', 'Correlation (r)', 'R²', 'Slope', 'Reading'])
stats_rows = [
    ('Sire show score vs offspring average (per sire)', f'=COUNT(Bloodline_Analysis!$J${r0}:$J${rS2})',
     f'=CORREL(Bloodline_Analysis!$J${r0}:$J${rS2},Bloodline_Analysis!$D${r0}:$D${rS2})',
     f'=RSQ(Bloodline_Analysis!$D${r0}:$D${rS2},Bloodline_Analysis!$J${r0}:$J${rS2})',
     f'=SLOPE(Bloodline_Analysis!$D${r0}:$D${rS2},Bloodline_Analysis!$J${r0}:$J${rS2})'),
    ('Dam-sire show score vs horse score', f'=COUNT({HSR("W")})', f'=CORREL({HSR("W")},{HSR("O")})',
     f'=RSQ({HSR("O")},{HSR("W")})', f'=SLOPE({HSR("O")},{HSR("W")})'),
    ('Mid-parent score vs horse score', f'=COUNT({HSR("X")})', f'=CORREL({HSR("X")},{HSR("O")})',
     f'=RSQ({HSR("O")},{HSR("X")})', f'=SLOPE({HSR("O")},{HSR("X")})'),
]
for i, rw in enumerate(stats_rows, start=6):
    for c, v in enumerate(rw, start=1): pp.cell(row=i, column=c, value=v)
    pp[f'F{i}'] = (f'=IF(ABS(C{i})>=0.7,"Strong",IF(ABS(C{i})>=0.4,"Moderate",IF(ABS(C{i})>=0.2,"Weak","Negligible")))'
                   f'&IF(B{i}<10," (few data points)","")')
body(pp, 6, 8, 1, 6, fmt={3: '+0.00;-0.00', 4: '0.00', 5: '+0.00;-0.00'})
pp['A9'] = 'Mid-parent score = average of sire and dam-sire show scores. Slope = offspring points gained per +1 parent point.'
pp['A9'].font = F(size=9, italic=True, color='8A8984')

section(pp, 'A11', '2 · Sire show score vs offspring')
header(pp, 12, 1, ['Sire', 'Sire show score', 'Offspring avg', 'Offspring vs sire', 'Offspring rank', 'Sire score rank'])
for i, s in enumerate(SIRES):
    r = 13 + i; src = r0 + i
    pp[f'A{r}'] = f'=Bloodline_Analysis!A{src}'
    pp[f'B{r}'] = f'=Bloodline_Analysis!J{src}'
    pp[f'C{r}'] = f'=Bloodline_Analysis!D{src}'
    pp[f'D{r}'] = f'=C{r}-B{r}'
    pp[f'E{r}'] = f'=RANK(C{r},$C$13:$C${12 + len(SIRES)},0)'
    pp[f'F{r}'] = f'=RANK(B{r},$B$13:$B${12 + len(SIRES)},0)'
    for c in 'ABC': pp[f'{c}{r}'].font = F(color='008000')
body(pp, 13, 12 + len(SIRES), 1, 6, fmt={2: '0.0', 3: '0.00', 4: '+0.00;-0.00'})
pp.conditional_formatting.add(f'D13:D{12 + len(SIRES)}', div_scale())

section(pp, 'A20', '3 · Horse-level parent scores (for the scatter chart)')
header(pp, 21, 1, ['HorseID', 'Registered name', 'Sire score', 'Dam-sire score', 'Mid-parent', 'Horse avg total', 'Vs sire'])
for i in range(NH):
    r = 22 + i; s = i + 2
    for c, col in zip('ABCDEFG', 'ABVWXOA'):
        pp[f'{c}{r}'] = f'=Horse_Summary!{col if c != "G" else "AA"}{s}'
        pp[f'{c}{r}'].font = F(color='008000')
body(pp, 22, 21 + NH, 1, 7, fmt={3: '0.0', 4: '0.0', 5: '0.00', 6: '0.00', 7: '+0.00;-0.00'})
widths(pp, {'A': 44, 'B': 24, 'C': 15, 'D': 15, 'E': 14, 'F': 26, 'G': 10})

sc = ScatterChart(); sc.title = 'Mid-parent score vs horse score'; sc.style = 13; sc.width = 15; sc.height = 8
ser = Series(Reference(pp, min_col=6, min_row=22, max_row=21 + NH), Reference(pp, min_col=5, min_row=22, max_row=21 + NH), title='Horses')
ser.marker = Marker(symbol='circle', size=7); ser.marker.graphicalProperties.solidFill = BLUE
ser.marker.graphicalProperties.line.solidFill = 'FFFFFF'; ser.graphicalProperties.line.noFill = True
from openpyxl.chart.trendline import Trendline
ser.trendline = Trendline(trendlineType='linear', dispRSqr=True)
sc.series.append(ser); sc.legend = None
sc.x_axis.title = 'Mid-parent show score'; sc.y_axis.title = 'Horse average total'
sc.x_axis.scaling.min = 107.5; sc.x_axis.scaling.max = 113; sc.y_axis.scaling.min = 102; sc.y_axis.scaling.max = 116
sc.x_axis.delete = False; sc.y_axis.delete = False
pp.add_chart(sc, 'I4')

sc2 = ScatterChart(); sc2.title = 'Sire show score vs offspring average'; sc2.style = 13; sc2.width = 15; sc2.height = 8
ser2 = Series(Reference(pp, min_col=3, min_row=13, max_row=12 + len(SIRES)), Reference(pp, min_col=2, min_row=13, max_row=12 + len(SIRES)), title='Sires')
ser2.marker = Marker(symbol='diamond', size=10); ser2.marker.graphicalProperties.solidFill = 'EB6834'
ser2.graphicalProperties.line.noFill = True
sc2.series.append(ser2); sc2.legend = None
sc2.x_axis.title = 'Sire show score'; sc2.y_axis.title = 'Offspring average'
sc2.x_axis.scaling.min = 111; sc2.x_axis.scaling.max = 113.5; sc2.y_axis.scaling.min = 104; sc2.y_axis.scaling.max = 114
sc2.x_axis.delete = False; sc2.y_axis.delete = False
pp.add_chart(sc2, 'I21')

# ======================= Age_Sex =======================
ag = sheet('Age_Sex', '1BAF7A')
title(ag, 'Age group & sex analysis', 'Average total score and horse count for each age group x sex combination.')
section(ag, 'A4', '1 · Average total score')
header(ag, 5, 1, ['Age group'] + SEXES + ['All sexes'])
for i, a in enumerate(AGES + ['All ages']):
    r = 6 + i; ag[f'A{r}'] = a
    for k, sx in enumerate(SEXES + ['All']):
        c = L(2 + k)
        if a == 'All ages' and sx == 'All': f = f'=AVERAGE({HSR("O")})'
        elif a == 'All ages': f = f'=IFERROR(AVERAGEIF({HSR("E")},{c}$5,{HSR("O")}),"")'
        elif sx == 'All': f = f'=IFERROR(AVERAGEIF({HSR("F")},$A{r},{HSR("O")}),"")'
        else: f = f'=IFERROR(AVERAGEIFS({HSR("O")},{HSR("F")},$A{r},{HSR("E")},{c}$5),"")'
        ag[f'{c}{r}'] = f
body(ag, 6, 9, 1, 7, fmt={k: '0.00' for k in range(2, 8)}, zebra=False)
ag['G5'] = 'All sexes'
ag.conditional_formatting.add('B6:F8', seq_scale(104, 113))
for c in range(1, 8): ag.cell(row=9, column=c).font = F(bold=True); ag.cell(row=9, column=c).fill = SUB_FILL

section(ag, 'A12', '2 · Horse count')
header(ag, 13, 1, ['Age group'] + SEXES + ['Total'])
for i, a in enumerate(AGES + ['Total']):
    r = 14 + i; ag[f'A{r}'] = a
    for k, sx in enumerate(SEXES + ['Total']):
        c = L(2 + k)
        if a == 'Total' and sx == 'Total': f = f'=COUNTA({HSR("A")})'
        elif a == 'Total': f = f'=COUNTIF({HSR("E")},{c}$13)'
        elif sx == 'Total': f = f'=COUNTIF({HSR("F")},$A{r})'
        else: f = f'=COUNTIFS({HSR("F")},$A{r},{HSR("E")},{c}$13)'
        ag[f'{c}{r}'] = f
body(ag, 14, 17, 1, 7, fmt={k: '0' for k in range(2, 8)}, zebra=False)
for c in range(1, 8): ag.cell(row=17, column=c).font = F(bold=True); ag.cell(row=17, column=c).fill = SUB_FILL

section(ag, 'A20', '3 · Trait averages by age group')
header(ag, 21, 1, ['Age group'] + TRAITS + ['Elite %'])
for i, a in enumerate(AGES):
    r = 22 + i; ag[f'A{r}'] = a
    for k, t in enumerate(TRAITS):
        ag.cell(row=r, column=2 + k, value=f'=AVERAGEIF({HSR("F")},$A{r},{HSR(TRAIT_COL[t])})')
    ag[f'H{r}'] = f'=COUNTIFS({HSR("F")},$A{r},{HSR("S")},"Elite")/COUNTIF({HSR("F")},$A{r})'
body(ag, 22, 24, 1, 8, fmt={**{k: '0.00' for k in range(2, 8)}, 8: '0%'})

section(ag, 'A27', '4 · Trait averages by sex')
header(ag, 28, 1, ['Sex'] + TRAITS + ['Elite %'])
for i, sx in enumerate(SEXES):
    r = 29 + i; ag[f'A{r}'] = sx
    for k, t in enumerate(TRAITS):
        ag.cell(row=r, column=2 + k, value=f'=AVERAGEIF({HSR("E")},$A{r},{HSR(TRAIT_COL[t])})')
    ag[f'H{r}'] = f'=COUNTIFS({HSR("E")},$A{r},{HSR("S")},"Elite")/COUNTIF({HSR("E")},$A{r})'
body(ag, 29, 33, 1, 8, fmt={**{k: '0.00' for k in range(2, 8)}, 8: '0%'})
ag['A35'] = ('Caution: cells hold 1–5 horses, and Horse_Profile flags sex labels that conflict with age '
             f'(=COUNTIF(Horse_Profile!Q:Q,"Sex/age mismatch") horses). Treat sex-class gaps as indicative.')
ag['A35'].font = F(size=9, italic=True, color='8A8984')
ag['A36'] = '="Horses with a sex/age mismatch flag: "&COUNTIF(Horse_Profile!$Q$2:$Q$' + str(HL) + ',"Sex/age mismatch")'
ag['A36'].font = F(size=9, bold=True, color='C0392B')
widths(ag, {'A': 16, **{L(i): 13 for i in range(2, 9)}})

ch = BarChart(); ch.type = 'col'; ch.grouping = 'clustered'
ch.add_data(Reference(ag, min_col=2, max_col=6, min_row=5, max_row=8), titles_from_data=True)
ch.set_categories(Reference(ag, min_col=1, min_row=6, max_row=8))
ch.title = 'Average total by age group and sex'; ch.width = 17; ch.height = 8; ch.style = 10
ch.y_axis.scaling.min = 100; ch.y_axis.scaling.max = 115; ch.legend.position = 'b'
ch.y_axis.delete = False; ch.x_axis.delete = False
for s, col in zip(ch.series, ['2A78D6', 'EB6834', '1BAF7A', 'EDA100', 'E87BA4']):
    s.graphicalProperties.solidFill = col; s.graphicalProperties.line.solidFill = col
ag.add_chart(ch, 'J4')

# ======================= Conformation =======================
cf = sheet('Conformation', '1BAF7A')
title(cf, 'Conformation strengths & weaknesses', 'Trait-level view across the herd. "Strongest/Weakest for" counts horses whose best/worst trait it is.')
header(cf, 4, 1, ['Trait', 'Herd average', 'Std dev', 'Min', 'Max', 'Strongest for (horses)', 'Weakest for (horses)', 'Net', 'Rank (avg)'])
for i, t in enumerate(TRAITS):
    r = 5 + i; c = TRAIT_COL[t]
    cf[f'A{r}'] = t
    cf[f'B{r}'] = f'=AVERAGE({HSR(c)})'
    cf[f'C{r}'] = f'=STDEV({HSR(c)})'
    cf[f'D{r}'] = f'=MIN({HSR(c)})'
    cf[f'E{r}'] = f'=MAX({HSR(c)})'
    cf[f'F{r}'] = f'=COUNTIF({HSR("T")},A{r})'
    cf[f'G{r}'] = f'=COUNTIF({HSR("U")},A{r})'
    cf[f'H{r}'] = f'=F{r}-G{r}'
    cf[f'I{r}'] = f'=RANK(B{r},$B$5:$B$10,0)'
body(cf, 5, 10, 1, 9, fmt={2: '0.00', 3: '0.00', 4: '0.00', 5: '0.00', 8: '+0;-0;0'})
cf.conditional_formatting.add('H5:H10', ColorScaleRule(start_type='num', start_value=-9, start_color='E34948',
                                                        mid_type='num', mid_value=0, mid_color='F7F6F3',
                                                        end_type='num', end_value=9, end_color='2A78D6'))
cf['A12'] = '="Herd strongest trait: "&INDEX(A5:A10,MATCH(MAX(B5:B10),B5:B10,0))&"   |   Herd weakest trait: "&INDEX(A5:A10,MATCH(MIN(B5:B10),B5:B10,0))'
cf['A12'].font = F(bold=True, color=NAVY)

section(cf, 'A14', 'Horses needing attention per trait (trait average below threshold)')
cf['A15'] = 'Threshold'; cf['B15'] = 17.5; cf['B15'].font = F(color='0000FF', bold=True); cf['B15'].fill = INPUT_FILL
cf['C15'] = '← input: change to re-count'; cf['C15'].font = F(size=9, italic=True, color='8A8984')
header(cf, 16, 1, ['Trait', 'Horses below threshold', '% of herd'])
for i, t in enumerate(TRAITS):
    r = 17 + i
    cf[f'A{r}'] = t
    cf[f'B{r}'] = f'=COUNTIF({HSR(TRAIT_COL[t])},"<"&$B$15)'
    cf[f'C{r}'] = f'=B{r}/COUNTA({HSR("A")})'
body(cf, 17, 22, 1, 3, fmt={3: '0%'})

section(cf, 'A25', 'Trait correlation matrix (horse-level averages)')
header(cf, 26, 1, [''] + TRAITS)
for i, t1 in enumerate(TRAITS):
    r = 27 + i; cf[f'A{r}'] = t1; cf[f'A{r}'].font = F(bold=True)
    for k, t2 in enumerate(TRAITS):
        cf.cell(row=r, column=2 + k, value=f'=CORREL({HSR(TRAIT_COL[t1])},{HSR(TRAIT_COL[t2])})')
body(cf, 27, 32, 1, 7, fmt={k: '0.00' for k in range(2, 8)}, zebra=False)
cf.conditional_formatting.add('B27:G32', ColorScaleRule(start_type='num', start_value=-0.6, start_color='E34948',
                                                         mid_type='num', mid_value=0, mid_color='F7F6F3',
                                                         end_type='num', end_value=1, end_color='2A78D6'))
widths(cf, {'A': 16, 'B': 16, 'C': 12, 'D': 10, 'E': 10, 'F': 20, 'G': 20, 'H': 8, 'I': 10})

ch = BarChart(); ch.type = 'bar'; ch.grouping = 'clustered'
ch.add_data(Reference(cf, min_col=6, max_col=7, min_row=4, max_row=10), titles_from_data=True)
ch.set_categories(Reference(cf, min_col=1, min_row=5, max_row=10))
ch.title = "Horses whose strongest / weakest trait is…"; ch.width = 16; ch.height = 8; ch.style = 10
ch.x_axis.scaling.orientation = 'maxMin'; ch.legend.position = 'b'; ch.y_axis.delete = False; ch.x_axis.delete = False
ch.series[0].graphicalProperties.solidFill = BLUE; ch.series[0].graphicalProperties.line.solidFill = BLUE
ch.series[1].graphicalProperties.solidFill = 'E34948'; ch.series[1].graphicalProperties.line.solidFill = 'E34948'
cf.add_chart(ch, 'K4')

# ======================= Judge_Analysis =======================
ja = sheet('Judge_Analysis', '4A3AA7')
title(ja, 'Judge consistency', 'Are the three judges scoring at the same level, and do they agree on which horses are best?')
header(ja, 4, 1, ['Judge'] + TRAITS + ['Avg total', 'Vs panel', 'Rows scored'])
for i, j in enumerate(JUDGES + ['Panel']):
    r = 5 + i; ja[f'A{r}'] = j
    for k, col in enumerate('EFGHIJK'):
        ja.cell(row=r, column=2 + k, value=(f'=AVERAGE(Judge_Scores!${col}$2:${col}${SL})' if j == 'Panel'
                                            else f'=AVERAGEIF(Judge_Scores!$D$2:$D${SL},$A{r},Judge_Scores!${col}$2:${col}${SL})'))
    ja[f'I{r}'] = f'=H{r}-$H$8'
    ja[f'J{r}'] = (f'=COUNTA(Judge_Scores!$A$2:$A${SL})' if j == 'Panel' else f'=COUNTIF(Judge_Scores!$D$2:$D${SL},A{r})')
body(ja, 5, 8, 1, 10, fmt={**{k: '0.00' for k in range(2, 9)}, 9: '+0.00;-0.00;0.00'})
for c in range(1, 11): ja.cell(row=8, column=c).font = F(bold=True); ja.cell(row=8, column=c).fill = SUB_FILL

section(ja, 'A11', 'Difference from panel average, by trait')
header(ja, 12, 1, ['Judge'] + TRAITS)
for i, j in enumerate(JUDGES):
    r = 13 + i; ja[f'A{r}'] = j
    for k in range(6):
        cl = L(2 + k); ja[f'{cl}{r}'] = f'={cl}{5 + i}-{cl}$8'
body(ja, 13, 15, 1, 7, fmt={k: '+0.00;-0.00;0.00' for k in range(2, 8)}, zebra=False)
ja.conditional_formatting.add('B13:G15', ColorScaleRule(start_type='num', start_value=-0.4, start_color='E34948',
                                                         mid_type='num', mid_value=0, mid_color='F7F6F3',
                                                         end_type='num', end_value=0.4, end_color='2A78D6'))

section(ja, 'A18', 'Agreement between judges (total score)')
header(ja, 19, 1, ['Pair', 'Correlation (r)', 'Reading'])
pairs = [('Judge A', 'Judge B', 'B', 'C'), ('Judge A', 'Judge C', 'B', 'D'), ('Judge B', 'Judge C', 'C', 'D')]
W0 = 30  # wide table start row
for i, (a, b, ca, cb) in enumerate(pairs):
    r = 20 + i
    ja[f'A{r}'] = f'{a} vs {b}'
    ja[f'B{r}'] = f'=CORREL(${ca}${W0 + 1}:${ca}${W0 + NH},${cb}${W0 + 1}:${cb}${W0 + NH})'
    ja[f'C{r}'] = f'=IF(B{r}>=0.7,"Strong",IF(B{r}>=0.4,"Moderate","Weak"))'
ja['A23'] = 'Average judge spread per horse (pts)'
ja['B23'] = f'=AVERAGE({HSR("P")})'
ja['A24'] = 'Largest judge spread (pts)'
ja['B24'] = f'=MAX({HSR("P")})'
ja['A25'] = 'Horse with largest spread'
ja['B25'] = f'=INDEX({HSR("B")},MATCH(B24,{HSR("P")},0))'
body(ja, 20, 25, 1, 3, fmt={2: '0.00'})

section(ja, f'A{W0 - 1}', 'Helper: each judge\'s total per horse')
header(ja, W0, 1, ['HorseID'] + JUDGES + ['Spread'])
for i in range(NH):
    r = W0 + 1 + i
    ja[f'A{r}'] = f'=Horse_Summary!A{i + 2}'
    for k, j in enumerate(JUDGES):
        ja.cell(row=r, column=2 + k, value=f'=SUMIFS(Judge_Scores!$K$2:$K${SL},Judge_Scores!$B$2:$B${SL},$A{r},Judge_Scores!$D$2:$D${SL},"{j}")')
    ja[f'E{r}'] = f'=MAX(B{r}:D{r})-MIN(B{r}:D{r})'
body(ja, W0 + 1, W0 + NH, 1, 5, fmt={k: '0.0' for k in range(2, 6)})
ja.conditional_formatting.add(f'E{W0 + 1}:E{W0 + NH}', seq_scale(0, 10))
widths(ja, {'A': 34, **{L(i): 13 for i in range(2, 11)}})

ch = BarChart(); ch.type = 'col'; ch.grouping = 'clustered'
ch.add_data(Reference(ja, min_col=2, max_col=7, min_row=12, max_row=15), titles_from_data=True)
ch.set_categories(Reference(ja, min_col=1, min_row=13, max_row=15))
ch.title = 'Judge scoring vs panel, by trait'; ch.width = 17; ch.height = 8; ch.style = 10; ch.legend.position = 'b'
ch.y_axis.number_format = '+0.00;-0.00;0'; ch.x_axis.tickLblPos = 'low'; ch.y_axis.delete = False; ch.x_axis.delete = False
for s, col in zip(ch.series, ['2A78D6', 'EB6834', '1BAF7A', 'EDA100', 'E87BA4', '4A3AA7']):
    s.graphicalProperties.solidFill = col; s.graphicalProperties.line.solidFill = col
ja.add_chart(ch, 'L4')

# ======================= Dashboard =======================
db = wb.create_sheet('Dashboard', 0)
db.sheet_view.showGridLines = False; db.sheet_properties.tabColor = 'C0392B'
db['B1'] = 'Arabian Horse Scoring & Bloodline Dashboard'; db['B1'].font = F(size=20, bold=True, color=NAVY)
db['B2'] = '40 horses · 3 judges · 6 AHA-style halter categories (0–20) · synthetic practice data — not official AHA results'
db['B2'].font = F(size=10, italic=True, color='52514E')
db.row_dimensions[1].height = 32
widths(db, {'A': 2, **{L(i): 13 for i in range(2, 20)}})

kpis = [('Horses', f'=COUNTA({HSR("A")})', '0'),
        ('Avg total score', f'=AVERAGE({HSR("O")})', '0.0'),
        ('Elite horses', f'=COUNTIF({HSR("S")},"Elite")', '0'),
        ('Top sire line', f'=INDEX(Bloodline_Analysis!$A${r0}:$A${rS2},MATCH(1,Bloodline_Analysis!$M${r0}:$M${rS2},0))', '@'),
        ('Herd weakest trait', '=INDEX(Conformation!$A$5:$A$10,MATCH(MIN(Conformation!$B$5:$B$10),Conformation!$B$5:$B$10,0))', '@'),
        ('Avg judge spread', f'=AVERAGE({HSR("P")})', '0.0 "pts"')]
for i, (lab, f, nf) in enumerate(kpis):
    c1 = 2 + i * 3; c2 = c1 + 2
    db.merge_cells(start_row=4, start_column=c1, end_row=4, end_column=c2)
    db.merge_cells(start_row=5, start_column=c1, end_row=6, end_column=c2)
    a = db.cell(row=4, column=c1, value=lab); a.font = F(size=9, bold=True, color='52514E'); a.alignment = CENTER
    v = db.cell(row=5, column=c1, value=f); v.font = F(size=18 if nf != '@' else 13, bold=True, color=NAVY)
    v.alignment = CENTER; v.number_format = nf
    for r in (4, 5, 6):
        for c in range(c1, c2 + 1):
            db.cell(row=r, column=c).fill = SUB_FILL
            db.cell(row=r, column=c).border = Border(top=Side(style='medium', color=BLUE) if r == 4 else None)

# Top 10
section(db, 'B8', 'Top 10 horses')
header(db, 9, 2, ['#', 'Horse', '', 'Sire', '', 'Sex', 'Age group', 'Avg total', 'Tier'])
db.merge_cells('C9:D9'); db.merge_cells('E9:F9')
for n in range(1, 11):
    r = 9 + n
    m = f'MATCH({n},{HSR("R")},0)'
    db[f'B{r}'] = n
    db[f'C{r}'] = f'=INDEX({HSR("B")},{m})'; db.merge_cells(f'C{r}:D{r}')
    db[f'E{r}'] = f'=INDEX({HSR("C")},{m})'; db.merge_cells(f'E{r}:F{r}')
    db[f'G{r}'] = f'=INDEX({HSR("E")},{m})'
    db[f'H{r}'] = f'=INDEX({HSR("F")},{m})'
    db[f'I{r}'] = f'=INDEX({HSR("O")},{m})'
    db[f'J{r}'] = f'=INDEX({HSR("S")},{m})'
body(db, 10, 19, 2, 10, fmt={9: '0.00'})
db.conditional_formatting.add('I10:I19', seq_scale(110, 115))

# Sire summary (linked)
section(db, 'L8', 'Sire lines at a glance')
header(db, 9, 12, ['Sire', '', 'Horses', 'Avg total', 'Elite %', 'Best trait', 'Weakest trait'])
db.merge_cells('L9:M9')
for i in range(len(SIRES)):
    r = 10 + i
    db[f'L{r}'] = f'=Bloodline_Analysis!A{r0 + i}'; db.merge_cells(f'L{r}:M{r}')
    db[f'N{r}'] = f'=Bloodline_Analysis!C{r0 + i}'
    db[f'O{r}'] = f'=Bloodline_Analysis!D{r0 + i}'
    db[f'P{r}'] = f'=Bloodline_Analysis!I{r0 + i}'
    db[f'Q{r}'] = f'=Bloodline_Analysis!H{t0 + 1 + i}'
    db[f'R{r}'] = f'=Bloodline_Analysis!I{t0 + 1 + i}'
body(db, 10, 9 + len(SIRES), 12, 18, fmt={15: '0.00', 16: '0%'})
db.conditional_formatting.add(f'O10:O{9 + len(SIRES)}', seq_scale(105, 113))

# Horse lookup
section(db, 'L17', 'Horse lookup — pick a horse from the yellow cell')
db['L18'] = 'Horse:'; db['L18'].font = F(bold=True)
db.merge_cells('M18:P18')
db['M18'] = profile.RegisteredName.iloc[0]
db['M18'].font = F(color='0000FF', bold=True); db['M18'].fill = INPUT_FILL
for c in range(13, 17): db.cell(row=18, column=c).fill = INPUT_FILL
dv = DataValidation(type='list', formula1=f'=Horse_Summary!$B$2:$B${HL}', allow_blank=False)
db.add_data_validation(dv); dv.add('M18')
LK = f'MATCH($M$18,{HSR("B")},0)'
info = [('Sire', 'C'), ('Dam sire', 'D'), ('Sex', 'E'), ('Age group', 'F'), ('Avg total', 'O'), ('Rank', 'Q'),
        ('Tier', 'S'), ('Strongest', 'T'), ('Focus area', 'U'), ('Judge spread', 'P')]
for i, (lab, col) in enumerate(info):
    r = 19 + i % 5; c = 12 + (i // 5) * 3
    db.cell(row=r, column=c, value=lab).font = F(color='52514E')
    v = db.cell(row=r, column=c + 1, value=f'=INDEX({HSR(col)},{LK})'); v.font = F(bold=True)
    if col in ('O', 'P'): v.number_format = '0.00'
header(db, 25, 12, ['Trait', 'Horse', 'Herd avg', 'Sire avg', 'Vs herd'])
for i, t in enumerate(TRAITS):
    r = 26 + i
    db[f'L{r}'] = t
    db[f'M{r}'] = f'=INDEX({HSR(TRAIT_COL[t])},{LK})'
    db[f'N{r}'] = f'=AVERAGE({HSR(TRAIT_COL[t])})'
    db[f'O{r}'] = f'=AVERAGEIF({HSR("C")},INDEX({HSR("C")},{LK}),{HSR(TRAIT_COL[t])})'
    db[f'P{r}'] = f'=M{r}-N{r}'
body(db, 26, 31, 12, 16, fmt={13: '0.00', 14: '0.00', 15: '0.00', 16: '+0.00;-0.00;0.00'})
db.conditional_formatting.add('P26:P31', ColorScaleRule(start_type='num', start_value=-1.5, start_color='E34948',
                                                         mid_type='num', mid_value=0, mid_color='F7F6F3',
                                                         end_type='num', end_value=1.5, end_color='2A78D6'))

# Dashboard charts
ch = style_chart(BarChart(), 'Average offspring score by sire', 14, 7)
ch.type = 'bar'; ch.add_data(Reference(db, min_col=15, min_row=9, max_row=9 + len(SIRES)), titles_from_data=True)
ch.set_categories(Reference(db, min_col=12, min_row=10, max_row=9 + len(SIRES)))
ch.y_axis.scaling.min = 100; ch.y_axis.scaling.max = 115; ch.x_axis.scaling.orientation = 'maxMin'
ch.dataLabels = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False, showPercent=False); ch.dataLabels.numFmt = '0.0'
ch.y_axis.delete = False; ch.x_axis.delete = False; ch.y_axis.majorGridlines = None
color_bars(ch, BLUE); db.add_chart(ch, 'B21')

ch = style_chart(BarChart(), 'Top 10 horses — average total', 14, 7)
ch.type = 'bar'; ch.add_data(Reference(db, min_col=9, min_row=9, max_row=19), titles_from_data=True)
ch.set_categories(Reference(db, min_col=3, min_row=10, max_row=19))
ch.y_axis.scaling.min = 108; ch.y_axis.scaling.max = 115; ch.x_axis.scaling.orientation = 'maxMin'
ch.dataLabels = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False, showPercent=False); ch.dataLabels.numFmt = '0.0'
ch.y_axis.delete = False; ch.x_axis.delete = False; ch.y_axis.majorGridlines = None
color_bars(ch, '1C5CAB'); db.add_chart(ch, 'B36')

ch = BarChart(); ch.type = 'col'; ch.grouping = 'clustered'
ch.add_data(Reference(db, min_col=13, max_col=14, min_row=25, max_row=31), titles_from_data=True)
ch.set_categories(Reference(db, min_col=12, min_row=26, max_row=31))
ch.title = 'Selected horse vs herd, by trait'; ch.width = 14; ch.height = 7; ch.style = 10; ch.legend.position = 'b'
ch.y_axis.scaling.min = 15; ch.y_axis.scaling.max = 20; ch.y_axis.delete = False; ch.x_axis.delete = False
ch.series[0].graphicalProperties.solidFill = BLUE; ch.series[0].graphicalProperties.line.solidFill = BLUE
ch.series[1].graphicalProperties.solidFill = 'B4B2A9'; ch.series[1].graphicalProperties.line.solidFill = 'B4B2A9'
db.add_chart(ch, 'L33')

ch = BarChart(); ch.type = 'bar'; ch.grouping = 'clustered'
ch.add_data(Reference(cf, min_col=6, max_col=7, min_row=4, max_row=10), titles_from_data=True)
ch.set_categories(Reference(cf, min_col=1, min_row=5, max_row=10))
ch.title = 'Strongest vs weakest trait (horse counts)'; ch.width = 14; ch.height = 7; ch.style = 10
ch.x_axis.scaling.orientation = 'maxMin'; ch.legend.position = 'b'; ch.y_axis.delete = False; ch.x_axis.delete = False
ch.series[0].graphicalProperties.solidFill = BLUE; ch.series[0].graphicalProperties.line.solidFill = BLUE
ch.series[1].graphicalProperties.solidFill = 'E34948'; ch.series[1].graphicalProperties.line.solidFill = 'E34948'
db.add_chart(ch, 'L48')

# ======================= About =======================
ab = wb.create_sheet('About', 1)
ab.sheet_view.showGridLines = False
title(ab, 'About this workbook')
lines = [
    ('Project', 'Horse Scoring & Bloodline Analytics — Excel component'),
    ('Data', 'Synthetic practice dataset: 40 Arabian horses, 120 judge score rows (3 judges x 40), 15 parent records. Only sire names are real; all scores and pedigrees are hypothetical.'),
    ('Scoring', 'Six AHA-style halter categories (Type, Head/Neck, Body/Topline, Legs, Movement, Balance), 0–20 each, total out of 120 per judge. Horse score = average of the 3 judges.'),
    ('', ''),
    ('Sheet guide', ''),
    ('Dashboard', 'KPIs, Top 10, sire summary, interactive horse lookup (yellow dropdown) and charts.'),
    ('Horse_Profile', 'Source profiles + computed Age and a sex/age data-quality flag.'),
    ('Judge_Scores', 'Source judge scores; TotalScore = SUM of the six traits.'),
    ('Parent_Reference', 'Sire / dam-sire reference + offspring count, average and gap vs parent score.'),
    ('Horse_Summary', 'One row per horse: trait averages, total, rank, tier, strength/focus, parent scores. Feeds all analysis sheets.'),
    ('Bloodline_Analysis', 'Sire line, dam sire, dam line and farm group comparisons; trait profile per sire.'),
    ('Parent_Performance', 'Correlation / slope of parent show scores vs offspring scores.'),
    ('Age_Sex', 'Age group x sex matrices for score, count and traits.'),
    ('Conformation', 'Herd trait averages, strongest/weakest counts, below-threshold counts, trait correlations.'),
    ('Judge_Analysis', 'Judge bias by trait and inter-judge agreement.'),
    ('Settings', 'Inputs: reference year, tier cut-offs, colt/filly age limit.'),
    ('', ''),
    ('Colour legend', 'Blue text on yellow = input you can change · Black = formula · Green = link to another sheet · Blue/red cell shading = above/below average.'),
    ('Formulas', 'Every analysis number is a live formula; edit Judge_Scores or Settings and everything updates.'),
]
for i, (a, b) in enumerate(lines, start=3):
    ab[f'A{i}'] = a; ab[f'A{i}'].font = F(bold=True, color=NAVY)
    ab[f'B{i}'] = b; ab[f'B{i}'].font = F(); ab[f'B{i}'].alignment = Alignment(wrap_text=True, vertical='top')
widths(ab, {'A': 22, 'B': 110})

# Fonts: make sure every populated cell is Arial
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and c.font.name != 'Arial':
                from copy import copy as _c; fnt=_c(c.font); fnt.name='Arial'; c.font=fnt

wb.save(OUT)
print('saved', OUT)
